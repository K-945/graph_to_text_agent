import os
from Graph import Graph
from langchain_community.document_loaders import UnstructuredFileLoader, PyPDFLoader
from langchain.text_splitter import CharacterTextSplitter, RecursiveCharacterTextSplitter
from langchain.embeddings.huggingface import HuggingFaceEmbeddings
from langchain_community.vectorstores import FAISS
from openpyxl import Workbook, load_workbook
from chatgpt_api import chatgpt_api, get_content
import json
import math
import re
import Logger as lg
import tiktoken
from langchain.docstore.document import Document
import nltk


def is_arabic_numeric(s_list):
    for s in s_list:
        if not s.isdigit():
            return False
    return True

def cnt_tokens(text):
    encoding = tiktoken.encoding_for_model("gpt-3.5-turbo-0613")
    return len(encoding.encode(text))

def count_words(article):
    return len(nltk.word_tokenize(article))

def splite_character_token(text, separator,token_limit):
    try:
        temp_splited_file = text.split(separator)
    except Exception as e:
        print(e)
        return None

    tokens_list = [cnt_tokens(t) for t in temp_splited_file]
    splited_file = []
    temp_cnt = 0
    temp_chunk = ''
    for i, sentence in enumerate(temp_splited_file):
        if temp_cnt + tokens_list[i] <= token_limit:
            temp_chunk += sentence + '.'
            temp_cnt += tokens_list[i]
        else:
            temp_chunk = temp_chunk.strip().strip('\n')
            splited_file.append(temp_chunk)
            temp_cnt = 0
            temp_chunk = ''
    if temp_chunk != '':
        temp_chunk = temp_chunk.strip('').strip('\n')
        splited_file.append(temp_chunk)    
    if cnt_tokens(splited_file[-1]) <= 100:
        splited_file[-2] = splited_file[-2] + splited_file[-1]
        splited_file.pop(-1)

    return splited_file

def langchain_configuration(text, word_limit):
    text_splitter = RecursiveCharacterTextSplitter(separators=['.\n\n', '.'], chunk_size=word_limit, chunk_overlap=0)
    splited_file = text_splitter.split_text(text)

    return splited_file

def langchain_configuration_file(file_path, word_limit):
    file_suffix = os.path.splitext(file_path)[1]
    text_content = ''
    if file_suffix == '.txt' or '.doc' or '.docx':
        loader = UnstructuredFileLoader(file_path)
        file_document = loader.load()
    elif file_suffix == '.pdf':
        loader = PyPDFLoader(file_path)
        file_document = loader.load_and_split()
    else:
        print('Please enter txt, doc, docx file!')
        return
    for content in file_document:
        text_content += content.page_content
    text_splitter = CharacterTextSplitter(separator="。", chunk_size=word_limit, chunk_overlap=50)
    splited_file = text_splitter.split_documents(file_document)

    return file_document, splited_file, text_content

def write_xlsx(excel_file_path, g):
    if not os.path.exists(excel_file_path):
        wb = Workbook()
        ws1 = wb.create_sheet('VertexInfo', 0)
        ws2 = wb.create_sheet('EdgeInfo',1)
        ws1['A1'], ws1['B1'], ws1['C1'], ws1['D1'] = \
            'id', 'content', 'summary', 'adjacency'
        ws2['A1'], ws2['B1'], ws2['C1'] = \
            'fromId', 'toId', 'weight'
    else:
        wb = load_workbook(excel_file_path)
        ws1 = wb['VertexInfo']
        ws2 = wb['EdgeInfo']

    vertexes_list = g.getVertexes()    
    row=2
    for vertex in vertexes_list:
        ws1.cell(row=row, column=1, value=vertex.id)
        ws1.cell(row=row, column=2, value=vertex.content)
        ws1.cell(row=row, column=3, value=vertex.summary)
        ws1.cell(row=row, column=4, value=json.dumps(vertex.adjacency))
        for to_vertex_id, weight in vertex.adjacency.items():
            ws2.cell(row=row, column=1, value=vertex.id)
            ws2.cell(row=row, column=2, value=to_vertex_id)
            ws2.cell(row=row, column=3, value=weight)
        row += 1
    wb.save(excel_file_path)

def write_vetex_xlsx(excel_file_path, g):
    if not os.path.exists(excel_file_path):
        wb = Workbook()
        ws1 = wb.create_sheet('VertexInfo', 0)
        ws2 = wb.create_sheet('EdgeInfo',1)
        ws1['A1'], ws1['B1'], ws1['C1'], ws1['D1'] = \
            'id', 'content', 'summary', 'adjacency'
        ws2['A1'], ws2['B1'], ws2['C1'] = \
            'fromId', 'toId', 'weight'
    else:
        wb = load_workbook(excel_file_path)
        ws1 = wb['VertexInfo']
        ws2 = wb['EdgeInfo']

    vertexes_list = g.getVertexes()
    row=2
    for vertex in vertexes_list:
        ws1.cell(row=row, column=1, value=vertex.id)
        ws1.cell(row=row, column=2, value=vertex.content)
        ws1.cell(row=row, column=3, value=vertex.summary)
        ws1.cell(row=row, column=4, value=json.dumps(vertex.adjacency))
        row += 1
    wb.save(excel_file_path)

def write_edge_xlsx(excel_file_path, g):
    if not os.path.exists(excel_file_path):
        wb = Workbook()
        ws1 = wb.create_sheet('VertexInfo', 0)
        ws2 = wb.create_sheet('EdgeInfo',1)
        ws1['A1'], ws1['B1'], ws1['C1'], ws1['D1'] = \
            'id', 'content', 'summary', 'adjacency'
        ws2['A1'], ws2['B1'], ws2['C1'] = \
            'fromId', 'toId', 'weight'
    else:
        wb = load_workbook(excel_file_path)
        ws1 = wb['VertexInfo']
        ws2 = wb['EdgeInfo']

    vertexes_list = g.getVertexes()
    row=2
    for vertex in vertexes_list:
        for to_vertex_id, weight in vertex.adjacency.items():
            ws2.cell(row=row, column=1, value=vertex.id)        
            ws2.cell(row=row, column=2, value=to_vertex_id)     
            ws2.cell(row=row, column=3, value=weight)
            row += 1
    wb.save(excel_file_path)

def read_xslx(excel_file_path, g):    
    wb = load_workbook(excel_file_path)
    ws1 = wb['VertexInfo']
    ws2 = wb['EdgeInfo']
    if g.numVertex == 0:
        for row in range(2, ws1.max_row + 1):
            id = ws1.cell(row=row, column=1).value
            content = ws1.cell(row=row, column=2).value
            summary = ws1.cell(row=row, column=3).value
            g.addVertex(int(id), content, summary)
    if g.getVertex(1).adjacency == {}:
        for row in range(2, ws2.max_row + 1):
            fromId = ws2.cell(row=row, column=1).value
            toId = ws2.cell(row=row, column=2).value
            weight = ws2.cell(row=row, column=3).value
            g.addEdge(int(fromId), int(toId), weight)
    return g

def construct_summary(splited_file, summary_file_path, excel_file_path, g):
    summary_limit = math.ceil(3000 / len(splited_file) - 1)
    if summary_limit > 1000:
        lower_limit = 800
    elif summary_limit > 500 and summary_limit < 800:
        lower_limit = 500
    else:
        lower_limit = 200
    article_summary = ''
    for id, text in enumerate(splited_file):
        text = text.page_content
        if cnt_tokens(text) > 100:
            print(f"The summary of the {id + 1}th: ")
            content = text
            prompt = f"""
            #########
            Paragraph:                  
            {content}
            #########
            
            You need to complete a summary task. 
            Read the above paragraphs to find the key information and summarize the following paragraph. The summary includes the main idea and the critical information as posible. 
            The summary must be between {lower_limit}-{summary_limit} words. Summarize strictly according to the content of the paragraphs provided.
            
            Summary:
            """
            tag = True
            while tag:
                try:
                    prompt = prompt.replace('            ', '').strip('\n')
                    summary = get_content(chatgpt_api(prompt.strip()))
                    tag = False
                except Exception as e:
                    continue
        else:
            summary = text
        summary = summary.replace("\n", "").replace("\t", "").replace("\r", "")
        summary = summary.strip()
        print(summary)
        chunk_summary = f'The {id + 1}th paragraph：' + summary + '\n'
        if id == 0:
            f = open(summary_file_path, 'w', encoding='utf-8')
            chunk_summary = chunk_summary.replace('\n', '   ') + '\n'
        else:
            f = open(summary_file_path, 'a', encoding='utf-8')
        f.write(chunk_summary)
        f.close()
        article_summary += chunk_summary
        g.addVertex(id + 1, content=content, summary=summary)
    write_vetex_xlsx(excel_file_path, g)

    return g, article_summary

def construct_relation(splited_file, relation_file_path, excel_file_path, g, article_summary, relation_word_limit):
    exists_relations = []
    if os.path.exists(relation_file_path):
        with open(relation_file_path, 'r', encoding='utf-8') as f:
            for line in f:
                exists_relations.append(line)
    exists_amount = len(exists_relations)
    exists_id = 0
    for i in range(len(splited_file)):
        j = i + 1
        while j < len(splited_file):
            if exists_id == exists_amount:
                print(f"The {i + 1}th paragraph and the {j + 1}th paragraph: ")
                prompt = f"""            
                You are a professional writer assigned with an article analysis task. 
                Read the article summary below and analyze the relationship between the {i+1}th paragraph and {j+1}th paragraph from the perspective of the entire article. Analyze the relationship between these two paragraphs strictly based on the provided article content.
                The analysis results should be {relation_word_limit} words max.
    
                #####
                Article Summary: 
                {article_summary}
                #####
    
                The Analysis results:
                """
                tag = True
                while tag:
                    try:
                        prompt = prompt.replace('                ', '').strip('\n')
                        relation = get_content(chatgpt_api(prompt.strip()))
                        tag = False
                    except Exception as e:
                        continue
                relation = relation.replace("\n", "").replace("\t", "").replace("\r", "")
                if count_words(relation) > 1000:
                    relation = compress_content(relation, False)
                print(relation)
                f = open(relation_file_path, 'a+', encoding='utf-8')
                f.write(f'The {i+1}th and the {j+1}th: ' + relation + '\n')
            else:
                relation = exists_relations[exists_id]
                exists_id += 1
                print(f'The relationship between {i+1}th and the {j+1}th already exists.')
            g.addEdge(i + 1, j + 1, relation)
            g.addEdge(j + 1, i + 1, relation)
            j += 1
    write_edge_xlsx(excel_file_path, g)
    return g

def search_text(article_summary, question, options):
    prompt = f"""    
    You are an information matching AI assiatant.
    You are provided an article, a question and options (maked by A, B, C, D). 
    Based on the analysis of the article and the question in the following, please determine which paragraphs in the article are relevant to the question and return the relevant paragraphs. Analyze as many correlations as possible between the information of paragraphs and question. Choose as many relevant paragraphs as you can, and select as many as are relevant.

    #####
    Article: 
    {article_summary}

    Question: 
    {question}
    
    Options:
    {options}
    #####
    
    If there are some relevant paragraphs was chosen, the numeric label of the relevant paragraph is returned, separated by commas when there are more than one paragraph. Output example: 1,3,7. If no relevant paragraph was chosen, output example: -1. 
    Returns only the numeric number of the relevant paragraph, do not return any option number and other information.
    """
    tag = True
    while tag:
        try:
            prompt = prompt.replace('    ', '').strip('\n')
            search_result = get_content(chatgpt_api(prompt.strip(), temperature=0.8))
            tag = False
        except Exception as e:
            continue
    print(f'These paragraphs are related to problem: {search_result}')
    return search_result

def analyze_text(g, question, search_paragraphs, num_search_paragraphs):
    notes = ''
    sum_query_times = (num_search_paragraphs * (num_search_paragraphs - 1)) // 2
    note_word_limit = math.ceil(3000 / sum_query_times)
    for i in range(num_search_paragraphs):
        j = i + 1
        while j < num_search_paragraphs:
            idx = int(search_paragraphs[i])
            jdx = int(search_paragraphs[j])
            content_i = g.vertexList[idx].content
            content_j = g.vertexList[jdx].content
            prompt = f"""
            You need to do a reading comprehension task. 
            The following are two paragraphs and the relationship between paragraphs in an essay, read the two paragraphs, the question and options, and return your analysis results. Analyze strictly according to the two paragraphs you have read.             

            #####
            The {idx}th paragraph:
            {content_i}

            The {jdx}th paragraph:
            {content_j}

            The relationship between the {idx}th paragraph and the {jdx}th paragraph:
            {g.vertexList[idx].adjacency[jdx]}

            Question:
            {question}
            #####

            Output Rule:
            The returned analysis result must be less than {note_word_limit} words.
            """
            tag = True
            while tag:
                try:
                    prompt = prompt.replace('            ', '').strip('\n')
                    stage_analysis = get_content(chatgpt_api(prompt.strip()))
                    tag = False
                except Exception as e:
                    continue
            print(f"Paragraphs {idx}th and {jdx}th are being analysed.")
            notes += stage_analysis
            if count_words(notes) > 2500:
                notes = compress_content(notes, True)
            j += 1

    return notes

def compress_content(content, is_note):
    if is_note:
        lower_limit = 1000
        upper_limit = 2500
    else:
        lower_limit = 700
        upper_limit = 1000
    prompt = f"""
    You need to complete a shorten article task. 
The shortened article must be between {lower_limit}-{upper_limit} words. Shorten strictly according to the content of the article provided in the following. The shortened article contains as much information as possible. 
    
    ########
    Artilce:                  
    {content}
    ########
    """

    tag = True
    while tag:
        try:
            prompt = prompt.replace('    ', '').strip('\n')
            result = get_content(chatgpt_api(prompt.strip()))
            tag = False
        except Exception as e:
            continue
    return result


def get_result(text, question, options):
    if count_words(text) >= 3800:
        text = compress_content(text, True)
    prompt = f"""
    #####
    You are a reading comprehension AI assistant. You are provided an article and a multiple-choice question with 4 possible answers (marked by A,B.C.D). Based on the content of the article, choose the best answer by writing its corresponding letter (either A,B,C,or D). Analyze the article and question thoroughly before selecting an option. 
    Only the letter number of the best answer is returned, do not provide any other information.

    #####
    Article：
    {text}

    Question：
    {question}

    Options:
    {options}
    #####
    
    Answer option:
    """
    tag = True
    while tag:
        try:
            prompt = prompt.replace('    ', '').strip('\n')
            result = get_content(chatgpt_api(prompt.strip(), temperature=0.3))
            tag = False
        except Exception as e:
            continue
    return result

def check_and_create_folder(folder_path):
    if not os.path.exists(folder_path):
        os.makedirs(folder_path)
        print(f'Folder {folder_path} created successfuly.')
    else:
        print(f'Folder {folder_path} already exists.')

def init_database(splited_file):
    embeddings = HuggingFaceEmbeddings(model_name='')
    database = FAISS.from_documents(splited_file, embedding=embeddings)
    return database

def retrieval_chunk(database, query, k=4):
    related_doc = database.similarity_search(query, k=k)
    search_result = ''
    for doc in related_doc:
        idx = re.match(r'The (\d{1,2})th paragraph:', doc.page_content[:50])[1]
        search_result += idx + ','
    return search_result[:-1]


if __name__ == '__main__':

    prefix_folder = 'result'
    if not os.path.exists(prefix_folder):
        os.makedirs(prefix_folder)

    accuracy_path = f'./{prefix_folder}/all_accuracy.txt'    
    
    
    file_path = 'paperqa.jsonl'    
    json_list = []

    all_questions = 0
    all_correct_number = 0
    all_accuracy = 0
    all_error_questions = []
    with open(file_path, 'r', encoding='utf-8') as f:
        for line in f:
            json_list.append(json.loads(line))

    for article in json_list:
        num = article['id']
        output_folder = f'./{prefix_folder}/GraphInformation_{num}'                
        file_name = 'paperqa'                           
        check_and_create_folder(output_folder)
        excel_file_path = os.path.join(output_folder , 'graph_' + file_name + '.xlsx')
        summary_file_path = os.path.join(output_folder , 'summary_' + file_name + '.txt')
        relation_file_path = os.path.join(output_folder , 'relation_' + file_name + '.txt')
        error_file_path = os.path.join(output_folder , 'error_' + file_name + '.txt')        
        splited_file = []
        print(f"===================== artilce number : {num} =====================")

        if len(article['introduction']) > 0:
            splited_file.append(article['introduction'])
        text = article['text']
        questions = article['questions']

        token_limit = 500
        separators = '.'
        splited_file = splite_character_token(text, separators, token_limit)

        for idx, chunk_text in enumerate(splited_file):
            chunk_text = chunk_text.replace('  ', '')
            splited_file[idx] = f'The {idx+1}th paragraph: {chunk_text}'
        splited_file = [Document(page_content=file) for file in splited_file]
        database = init_database(splited_file)
        print(f'The file is splited into {len(splited_file)} chunk.')

        g = Graph()
        if os.path.exists(summary_file_path):            
            article_summary = ''
            with open(summary_file_path, 'r', encoding='utf-8') as f:
                article_summary = f.read()
            print(f"The summary has already exists.")
            g = read_xslx(excel_file_path, g)
        else:
            g, article_summary = construct_summary(splited_file, summary_file_path, excel_file_path, g)
        
        if os.path.exists(relation_file_path):
            g = read_xslx(excel_file_path, g)
            print(f"The relation has already exists.")
        else:
            splited_file = splited_file
            relation_word_limit = math.ceil(3000 / len(splited_file) - 1)
            g = construct_relation(splited_file, relation_file_path, excel_file_path, g, article_summary,
                                   relation_word_limit)

        correct_number = 0
        accuracy = 0
        error_questions = []
        for i, qa in enumerate(questions):
            print(f'************ Start analyzing the {i}th question **********')
            question = qa['question']
            options = qa['options']
            gold_label = qa['gold_label']
            tag = True
            cannot_answer = 0
            not_arabic_numeric = 0
            while tag:
                question = qa['question']
                options = qa['options']
                gold_label = qa['gold_label']
                result = ''
                if g.getNumVertex() == 1:
                    result = get_result(g.getVertex(1).content, question, options)                
                else:
                    search_result = search_text(article_summary, question, options)
                    search_result = search_result.replace(' ', '')
                    print(search_result)
                    try:
                        search_paragraphs = search_result.split(',')
                        if not is_arabic_numeric(search_paragraphs):
                            not_arabic_numeric += 1
                        num_search_paragraphs = len(search_paragraphs)
                        if search_paragraphs[0] == '-1' or not_arabic_numeric > 5:
                            search_result = retrieval_chunk(database, question + options)
                            search_paragraphs = search_result.split(',')
                            num_search_paragraphs = len(search_paragraphs)
                            print(f'search_result: {search_result}')
                        if num_search_paragraphs == 0:
                            print(f'result: {result}')
                            if cannot_answer != 5:
                                cannot_answer += 1
                                continue
                            else:
                                search_result = retrieval_chunk(database, question + options)
                                search_paragraphs = search_result.split(',')
                        elif num_search_paragraphs == 1 and search_paragraphs[0] != '-1':
                            result = get_result(g.getVertex(int(search_paragraphs[0])).content, question, options)
                        else:
                            notes = analyze_text(g, question, search_paragraphs, num_search_paragraphs)
                            result = get_result(notes, question, options)
                    except Exception as e:
                        if result != '':
                            error_questions.append([f'number: {i}', question, result, f'gold_label: {gold_label}', f'The {num}th article, the {i}th question, gold_label: {gold_label}, current result: {result}'])
                if len(result) == 0:
                    print(f'result: {result}')
                    if cannot_answer != 10:
                        cannot_answer += 1
                        continue
                    else:
                        result = "-1"
                if result[0] == gold_label:
                    correct_number += 1
                    print(f'=== correct: {correct_number} ===')
                else:
                    error_questions.append([f'number: {i}', question, result, f'gold_label: {gold_label}', f'The {num}th article, the {i}th question, gold_label: {gold_label}, current result: {result}'])
                print(f'The {num}th article, the {i}th question, gold_label: {gold_label}, current result: {result}')
                tag = False        
        accuracy = correct_number / len(questions)
        print(f'The {num}th article, accuracy: {accuracy}')
        all_questions += len(questions)
        all_correct_number += correct_number
        all_error_questions.append(error_questions)

    all_accuracy = all_correct_number / all_questions
    print(f'all_questions: {all_questions}, all_correct_number: {all_correct_number} ,all accuracy: {all_accuracy}')
    with open(accuracy_path, 'a') as f:
        f.write(f'all_questions: {all_questions}, all_accuracy: {all_accuracy}\n\n')



